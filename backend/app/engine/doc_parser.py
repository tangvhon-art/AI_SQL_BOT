"""文件解析引擎：PDF/Word/Excel/CSV/Markdown/TXT/图片 → 带页码元数据的文本片段。
扫描版 PDF 和图片通过 OCR 识别文字。解析结果供 doc_store 切分向量化。"""
from __future__ import annotations

import logging
import os
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)

# 支持的格式白名单
SUPPORTED_EXTS = {"pdf", "docx", "xlsx", "csv", "md", "markdown", "txt", "png", "jpg", "jpeg"}
# 需要 OCR 的格式
OCR_EXTS = {"png", "jpg", "jpeg"}
# 文字版格式
TEXT_EXTS = {"md", "markdown", "txt", "csv"}

MAX_FILE_SIZE = 20 * 1024 * 1024  # 20MB
OCR_TEXT_THRESHOLD = 50  # PDF 总文字 < 50 字符判定为扫描版


@dataclass
class RawPage:
    """解析后的原始页/块，保留元数据供切分继承。"""
    text: str
    page_num: int | None = None
    sheet_name: str | None = None


@dataclass
class ParseResult:
    success: bool
    pages: list[RawPage] = field(default_factory=list)
    error: str | None = None

    @property
    def total_text(self) -> str:
        return "\n".join(p.text for p in self.pages if p.text.strip())

    @property
    def char_count(self) -> int:
        return sum(len(p.text) for p in self.pages)


# ---------- OCR 懒加载（首次使用才初始化，避免无 OCR 需求时加载模型） ----------
_ocr_engine = None
_ocr_init_failed = False


def _get_ocr():
    global _ocr_engine, _ocr_init_failed
    if _ocr_engine is not None:
        return _ocr_engine
    if _ocr_init_failed:
        return None
    try:
        from rapidocr_onnxruntime import RapidOCR
        _ocr_engine = RapidOCR()
        logger.info("OCR 引擎初始化成功 (rapidocr-onnxruntime)")
        return _ocr_engine
    except Exception as exc:  # noqa: BLE001
        _ocr_init_failed = True
        logger.warning("OCR 引擎初始化失败: %s", exc)
        return None


def _ocr_image_bytes(img_bytes: bytes) -> str:
    """对图片字节做 OCR，返回识别文字。"""
    engine = _get_ocr()
    if engine is None:
        return ""
    try:
        import numpy as np
        from PIL import Image
        import io
        img = Image.open(io.BytesIO(img_bytes)).convert("RGB")
        arr = np.array(img)
        result, _ = engine(arr)
        if not result:
            return ""
        # result: [[box, text, score], ...]，按行拼接
        lines = [item[1] for item in result if item and len(item) >= 2]
        return "\n".join(lines)
    except Exception as exc:  # noqa: BLE001
        logger.warning("OCR 识别失败: %s", exc)
        return ""


# ---------- 各格式解析器 ----------

def _parse_pdf(path: str) -> ParseResult:
    """PDF 解析：优先文字层，文字层为空则逐页 OCR。"""
    try:
        import pdfplumber
    except ImportError:
        return ParseResult(success=False, error="缺少 pdfplumber 依赖")

    pages: list[RawPage] = []
    total_text_len = 0
    try:
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                pages.append(RawPage(text=text, page_num=i))
                total_text_len += len(text.strip())
    except Exception as exc:  # noqa: BLE001
        return ParseResult(success=False, error=f"PDF 解析失败: {exc}")

    # 扫描版判定：总文字过少，逐页 OCR
    if total_text_len < OCR_TEXT_THRESHOLD:
        logger.info("PDF 文字层为空（%d 字符），触发 OCR", total_text_len)
        ocr_pages: list[RawPage] = []
        try:
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                for i, page in enumerate(pdf.pages, 1):
                    # 将 PDF 页渲染为图片后 OCR
                    try:
                        img = page.to_image(resolution=150)
                        img_bytes = img.original.convert("RGB")
                        import io
                        buf = io.BytesIO()
                        img_bytes.save(buf, format="PNG")
                        text = _ocr_image_bytes(buf.getvalue())
                    except Exception:  # noqa: BLE001
                        text = ""
                    ocr_pages.append(RawPage(text=text or pages[i - 1].text, page_num=i))
            pages = ocr_pages
        except Exception as exc:  # noqa: BLE001
            logger.warning("PDF OCR 失败，回退文字层: %s", exc)

    if not any(p.text.strip() for p in pages):
        return ParseResult(success=False, error="PDF 无有效文字内容（可能是加密或纯图片）")
    return ParseResult(success=True, pages=pages)


def _parse_docx(path: str) -> ParseResult:
    """Word 解析：段落 + 表格。"""
    try:
        from docx import Document
    except ImportError:
        return ParseResult(success=False, error="缺少 python-docx 依赖")
    try:
        doc = Document(path)
        parts: list[str] = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        if not text.strip():
            return ParseResult(success=False, error="Word 文档无有效文字内容")
        return ParseResult(success=True, pages=[RawPage(text=text, page_num=1)])
    except Exception as exc:  # noqa: BLE001
        return ParseResult(success=False, error=f"Word 解析失败: {exc}")


def _parse_xlsx(path: str) -> ParseResult:
    """Excel 解析：每个 sheet 转为文本表格表示。"""
    try:
        import pandas as pd
    except ImportError:
        return ParseResult(success=False, error="缺少 pandas 依赖")
    try:
        xls = pd.ExcelFile(path)
        pages: list[RawPage] = []
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name, dtype=str)
            if df.empty:
                continue
            # 填充 NaN
            df = df.fillna("")
            # 表头 + 行数据，按行块切（每块 50 行，避免单 chunk 过大）
            header = " | ".join(str(c) for c in df.columns)
            block_size = 50
            for start in range(0, len(df), block_size):
                block = df.iloc[start:start + block_size]
                lines = [header]
                for _, row in block.iterrows():
                    lines.append(" | ".join(str(v) for v in row.values))
                pages.append(RawPage(
                    text="\n".join(lines),
                    page_num=start // block_size + 1,
                    sheet_name=str(sheet_name),
                ))
        if not pages:
            return ParseResult(success=False, error="Excel 无有效数据")
        return ParseResult(success=True, pages=pages)
    except Exception as exc:  # noqa: BLE001
        return ParseResult(success=False, error=f"Excel 解析失败: {exc}")


def _parse_csv(path: str) -> ParseResult:
    """CSV 解析：自动检测编码。"""
    try:
        import pandas as pd
    except ImportError:
        return ParseResult(success=False, error="缺少 pandas 依赖")
    text = _read_text_with_encoding(path)
    if text is None:
        return ParseResult(success=False, error="CSV 编码检测失败")
    try:
        import io
        df = pd.read_csv(io.StringIO(text), dtype=str)
        df = df.fillna("")
        header = " | ".join(str(c) for c in df.columns)
        lines = [header]
        for _, row in df.iterrows():
            lines.append(" | ".join(str(v) for v in row.values))
        return ParseResult(success=True, pages=[RawPage(text="\n".join(lines), page_num=1)])
    except Exception as exc:  # noqa: BLE001
        return ParseResult(success=False, error=f"CSV 解析失败: {exc}")


def _read_text_with_encoding(path: str) -> str | None:
    """读取文本文件，自动检测编码（UTF-8/GBK/GB2312）。"""
    raw = None
    try:
        with open(path, "rb") as f:
            raw = f.read()
    except Exception:  # noqa: BLE001
        return None
    # 尝试 chardet
    try:
        import chardet
        det = chardet.detect(raw)
        enc = det.get("encoding")
        if enc:
            try:
                return raw.decode(enc, errors="replace")
            except Exception:  # noqa: BLE001
                pass
    except ImportError:
        pass
    # 回退依次尝试
    for enc in ("utf-8", "gbk", "gb2312", "utf-16"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return None


def _parse_text(path: str) -> ParseResult:
    """纯文本 / Markdown 解析。"""
    text = _read_text_with_encoding(path)
    if text is None:
        return ParseResult(success=False, error="文本文件编码检测失败")
    if not text.strip():
        return ParseResult(success=False, error="文件内容为空")
    return ParseResult(success=True, pages=[RawPage(text=text, page_num=1)])


def _parse_image(path: str) -> ParseResult:
    """图片 OCR 解析。"""
    try:
        with open(path, "rb") as f:
            img_bytes = f.read()
    except Exception as exc:  # noqa: BLE001
        return ParseResult(success=False, error=f"图片读取失败: {exc}")
    text = _ocr_image_bytes(img_bytes)
    if not text.strip():
        return ParseResult(success=False, error="图片 OCR 未识别到文字")
    return ParseResult(success=True, pages=[RawPage(text=text, page_num=1)])


# ---------- 统一入口 ----------

_PARSERS = {
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "xlsx": _parse_xlsx,
    "csv": _parse_csv,
    "md": _parse_text,
    "markdown": _parse_text,
    "txt": _parse_text,
    "png": _parse_image,
    "jpg": _parse_image,
    "jpeg": _parse_image,
}


def get_ext(file_name: str) -> str:
    return file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""


def parse_file(file_path: str, file_name: str) -> ParseResult:
    """统一文件解析入口。

    Args:
        file_path: 服务器上的临时文件路径
        file_name: 原始文件名（用于判断扩展名）

    Returns:
        ParseResult: 解析结果，pages 为带页码元数据的文本片段列表
    """
    ext = get_ext(file_name)
    if ext not in SUPPORTED_EXTS:
        return ParseResult(success=False, error=f"不支持的文件格式: .{ext}")

    # 文件大小校验
    try:
        size = os.path.getsize(file_path)
        if size > MAX_FILE_SIZE:
            return ParseResult(success=False, error=f"文件过大（{size // 1024 // 1024}MB），上限 20MB")
    except OSError:
        pass

    parser = _PARSERS.get(ext)
    if parser is None:
        return ParseResult(success=False, error=f"无解析器: .{ext}")

    logger.info("开始解析文件: %s (ext=%s)", file_name, ext)
    result = parser(file_path)
    if result.success:
        logger.info("文件解析成功: %s, %d 页, %d 字符",
                    file_name, len(result.pages), result.char_count)
    else:
        logger.warning("文件解析失败: %s, %s", file_name, result.error)
    return result
